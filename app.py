from flask import Flask, render_template, request, redirect, url_for, send_file, session, jsonify
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from functools import wraps
import sqlite3
import os
import csv
import io
import requests

app = Flask(__name__)
DB_PATH = os.path.join(os.path.dirname(__file__), "chamados.db")
FUSO_BRASIL = ZoneInfo("America/Sao_Paulo")

# --- Segurança (via variáveis de ambiente) ---
app.secret_key = os.environ.get("SECRET_KEY") or "dev-secret-mude-em-producao"
ADMIN_USER = os.environ.get("ADMIN_USER", "admin")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin")

# --- Telegram ---
TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN")
TELEGRAM_CHAT_ID_PADRAO = os.environ.get("TELEGRAM_CHAT_ID")
# Chat por setor (opcional). Se o setor não tiver chat, usa o padrão.
TELEGRAM_CHAT_IDS = {
    "TI": os.environ.get("TELEGRAM_CHAT_TI"),
    "RH": os.environ.get("TELEGRAM_CHAT_RH"),
    "Financeiro": os.environ.get("TELEGRAM_CHAT_FIN"),
    "Almoxarifado": os.environ.get("TELEGRAM_CHAT_ALMOX"),
    "Manutenção": os.environ.get("TELEGRAM_CHAT_MANUT"),
    "ADM": os.environ.get("TELEGRAM_CHAT_ADM"),
}

SETORES = ["TI", "Almoxarifado", "RH", "Financeiro", "ADM", "Manutenção"]
STATUS_OPCOES = ["Aberto", "Em andamento", "Concluído"]

CLASSIFICACAO = {
    "Almoxarifado": [
        "mouse", "teclado", "monitor", "cadeira", "mesa", "papel", "caneta",
        "grampeador", "toner", "cartucho", "headset", "fone", "mochila",
        "pasta", "caixa", "material de escritorio", "notebook", "cabo",
    ],
    "TI": [
        "senha", "internet", "rede", "wifi", "sistema", "software", "login",
        "acesso", "impressora", "computador travando", "email", "e-mail",
        "instalar", "formatar", "antivirus", "vpn", "servidor",
    ],
    "Manutenção": [
        "ar condicionado", "lampada", "lâmpada", "tomada", "encanamento",
        "vazamento", "porta", "janela", "eletrica", "elétrica", "infiltração",
    ],
    "RH": [
        "ferias", "férias", "salario", "salário", "holerite", "contracheque",
        "admissao", "admissão", "demissao", "demissão", "rescisao", "rescisão",
        "beneficio", "benefício", "vale transporte", "vale-transporte",
        "vale alimentacao", "vale-alimentacao", "atestado", "plano de saude",
        "plano de saúde", "ponto", "banco de horas", "advertencia", "advertência",
    ],
    "Financeiro": [
        "reembolso", "nota fiscal", "pagamento", "boleto", "fatura",
        "despesa", "adiantamento", "nota de debito", "centro de custo",
    ],
}


def classificar_setor(item_texto):
    texto = item_texto.lower()
    for setor, palavras in CLASSIFICACAO.items():
        if any(p in texto for p in palavras):
            return setor
    return "Não classificado"


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def agora_utc_str():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def formatar_data_br(data_iso):
    try:
        dt_utc = datetime.strptime(data_iso, "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
        dt_br = dt_utc.astimezone(FUSO_BRASIL)
        return dt_br.strftime("%d/%m/%Y %H:%M")
    except (ValueError, TypeError):
        return data_iso


app.jinja_env.filters["br_data"] = formatar_data_br


# --- Autenticação ---
def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logado"):
            return redirect(url_for("index"))
        return f(*args, **kwargs)
    return wrapper


@app.context_processor
def injetar_estado_login():
    return {
        "logado": bool(session.get("logado")),
        "usuario_nome": session.get("usuario", ""),
    }


@app.route("/api/login", methods=["POST"])
def api_login():
    dados = request.get_json(silent=True) or {}
    usuario = (dados.get("usuario") or "").strip()
    senha = dados.get("senha") or ""
    if usuario == ADMIN_USER and senha == ADMIN_PASSWORD:
        session["logado"] = True
        session["usuario"] = usuario
        return jsonify({"ok": True, "usuario": usuario})
    return jsonify({"ok": False, "erro": "credenciais_invalidas"}), 401


@app.route("/api/logout", methods=["POST"])
def api_logout():
    session.clear()
    return jsonify({"ok": True})


# --- Histórico ---
def registrar_historico(conn, chamado_id, campo, valor_antigo, valor_novo):
    conn.execute(
        """INSERT INTO historico
           (chamado_id, campo, valor_antigo, valor_novo, usuario, data)
           VALUES (?, ?, ?, ?, ?, ?)""",
        (chamado_id, campo, valor_antigo or "—", valor_novo or "—",
         session.get("usuario", "sistema"), agora_utc_str()),
    )


# --- Telegram ---
def enviar_notificacao_telegram(mensagem, setor=None):
    """Envia pro chat do setor responsável. Se não tiver, usa o padrão."""
    if not TELEGRAM_TOKEN:
        return
    chat_id = (TELEGRAM_CHAT_IDS.get(setor) if setor else None) or TELEGRAM_CHAT_ID_PADRAO
    if not chat_id:
        return
    url = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/sendMessage"
    payload = {
        "chat_id": chat_id,
        "text": mensagem,
        "parse_mode": "HTML",
        "disable_web_page_preview": True,
    }
    try:
        r = requests.post(url, json=payload, timeout=5)
        if not r.ok:
            print(f"[Telegram] Erro {r.status_code}: {r.text}")
    except Exception as e:
        print(f"[Telegram] Exceção: {e}")


def init_db():
    conn = get_db()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS chamados (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data_solicitacao TEXT NOT NULL,
            setor_solicitante TEXT NOT NULL,
            nome_solicitante TEXT NOT NULL,
            item_solicitado TEXT NOT NULL,
            setor_responsavel TEXT NOT NULL,
            nome_responsavel TEXT,
            status TEXT NOT NULL DEFAULT 'Aberto'
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS historico (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            chamado_id INTEGER NOT NULL,
            campo TEXT NOT NULL,
            valor_antigo TEXT,
            valor_novo TEXT,
            usuario TEXT NOT NULL,
            data TEXT NOT NULL
        )
    """)
    conn.commit()
    conn.close()


@app.route("/")
def index():
    conn = get_db()
    filtro_status = request.args.get("status", "")
    filtro_setor = request.args.get("setor_responsavel", "")
    ordem = request.args.get("ordem", "recentes")
    if ordem not in ("recentes", "antigos"):
        ordem = "recentes"

    query = "SELECT * FROM chamados WHERE 1=1"
    params = []
    if filtro_status:
        query += " AND status = ?"
        params.append(filtro_status)
    if filtro_setor:
        query += " AND setor_responsavel = ?"
        params.append(filtro_setor)

    query += " ORDER BY data_solicitacao " + ("DESC" if ordem == "recentes" else "ASC")
    proxima_ordem = "antigos" if ordem == "recentes" else "recentes"

    chamados = conn.execute(query, params).fetchall()
    conn.close()
    return render_template(
        "index.html",
        chamados=chamados,
        setores=SETORES,
        status_opcoes=STATUS_OPCOES,
        filtro_status=filtro_status,
        filtro_setor=filtro_setor,
        ordem=ordem,
        proxima_ordem=proxima_ordem,
    )


@app.route("/novo", methods=["GET", "POST"])
def novo_chamado():
    if request.method == "POST":
        item_solicitado = request.form["item_solicitado"]
        setor_responsavel = classificar_setor(item_solicitado)
        nome_solicitante = request.form["nome_solicitante"]
        setor_solicitante = request.form["setor_solicitante"]

        conn = get_db()
        cursor = conn.execute(
            """INSERT INTO chamados
               (data_solicitacao, setor_solicitante, nome_solicitante,
                item_solicitado, setor_responsavel, nome_responsavel, status)
               VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (agora_utc_str(), setor_solicitante, nome_solicitante,
             item_solicitado, setor_responsavel, "", "Aberto"),
        )
        novo_id = cursor.lastrowid
        conn.commit()
        conn.close()

        data_br = datetime.now(FUSO_BRASIL).strftime("%d/%m/%Y às %H:%M")
        mensagem = (
            f"🔔 <b>Novo chamado #{novo_id}</b>\n\n"
            f"📅 <b>Data:</b> {data_br}\n"
            f"👤 <b>Solicitante:</b> {nome_solicitante}\n"
            f"🏢 <b>Setor solicitante:</b> {setor_solicitante}\n"
            f"🎯 <b>Setor responsável:</b> {setor_responsavel}\n"
            f"📦 <b>Item:</b> {item_solicitado}"
        )
        # Manda pro Telegram do setor responsável (ou padrão como fallback)
        enviar_notificacao_telegram(mensagem, setor=setor_responsavel)

        return redirect(url_for("index"))

    return render_template("novo.html", setores=SETORES)


@app.route("/chamado/<int:chamado_id>/status", methods=["POST"])
@login_required
def atualizar_status(chamado_id):
    novo_status = request.form["status"]
    nome_responsavel = request.form.get("nome_responsavel", "")

    conn = get_db()
    chamado = conn.execute("SELECT * FROM chamados WHERE id = ?", (chamado_id,)).fetchone()
    if not chamado:
        conn.close()
        return redirect(url_for("index"))

    if chamado["status"] != novo_status:
        registrar_historico(conn, chamado_id, "status", chamado["status"], novo_status)
    if (chamado["nome_responsavel"] or "") != nome_responsavel:
        registrar_historico(conn, chamado_id, "responsável", chamado["nome_responsavel"], nome_responsavel)

    conn.execute(
        "UPDATE chamados SET status = ?, nome_responsavel = ? WHERE id = ?",
        (novo_status, nome_responsavel, chamado_id),
    )
    conn.commit()
    conn.close()
    return redirect(url_for("index"))


@app.route("/chamado/<int:chamado_id>/excluir", methods=["POST"])
@login_required
def excluir_chamado(chamado_id):
    conn = get_db()
    conn.execute("DELETE FROM historico WHERE chamado_id = ?", (chamado_id,))
    conn.execute("DELETE FROM chamados WHERE id = ?", (chamado_id,))
    conn.commit()
    conn.close()
    return redirect(url_for("index"))


@app.route("/api/chamado/<int:chamado_id>/historico")
def api_historico(chamado_id):
    conn = get_db()
    chamado = conn.execute("SELECT * FROM chamados WHERE id = ?", (chamado_id,)).fetchone()
    if not chamado:
        conn.close()
        return jsonify({"ok": False, "erro": "nao_encontrado"}), 404
    historico = conn.execute(
        "SELECT * FROM historico WHERE chamado_id = ? ORDER BY data DESC",
        (chamado_id,),
    ).fetchall()
    conn.close()

    return jsonify({
        "ok": True,
        "chamado": {
            "id": chamado["id"],
            "solicitante": chamado["nome_solicitante"],
            "setor_solicitante": chamado["setor_solicitante"],
            "item": chamado["item_solicitado"],
            "setor_responsavel": chamado["setor_responsavel"],
            "status": chamado["status"],
        },
        "historico": [{
            "campo": h["campo"],
            "de": h["valor_antigo"],
            "para": h["valor_novo"],
            "usuario": h["usuario"],
            "data": formatar_data_br(h["data"]),
        } for h in historico],
    })


# ============ RELATÓRIOS ============
@app.route("/relatorios")
def relatorios():
    return render_template("relatorios.html")


def _buscar_chamados_para_relatorio():
    conn = get_db()
    chamados = conn.execute("SELECT * FROM chamados ORDER BY data_solicitacao DESC").fetchall()
    conn.close()
    return chamados


def _dados_do_chamado(c):
    return {
        "id": c["id"],
        "data": formatar_data_br(c["data_solicitacao"]),
        "solicitante": c["nome_solicitante"],
        "setor_solicitante": c["setor_solicitante"],
        "item": c["item_solicitado"],
        "setor_responsavel": c["setor_responsavel"],
        "responsavel": c["nome_responsavel"] or "—",
        "status": c["status"],
    }


@app.route("/relatorios/gerar")
def gerar_relatorio():
    formato = request.args.get("formato", "csv").lower()
    if formato not in ("csv", "xlsx", "pdf"):
        formato = "csv"
    chamados = _buscar_chamados_para_relatorio()
    dados = [_dados_do_chamado(c) for c in chamados]
    data_geracao = datetime.now(FUSO_BRASIL).strftime("%d/%m/%Y às %H:%M")
    if formato == "csv":
        return _gerar_csv(dados, data_geracao)
    elif formato == "xlsx":
        return _gerar_xlsx(dados, data_geracao)
    return _gerar_pdf(dados, data_geracao)


def _gerar_csv(dados, data_geracao):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Relatório de Chamados"])
    writer.writerow([f"Gerado em: {data_geracao}"])
    writer.writerow([f"Total: {len(dados)} chamado(s)"])
    writer.writerow([])
    writer.writerow(["ID", "Data", "Solicitante", "Setor Solicitante",
                     "Item", "Setor Responsável", "Responsável", "Status"])
    for d in dados:
        writer.writerow([d["id"], d["data"], d["solicitante"], d["setor_solicitante"],
                         d["item"], d["setor_responsavel"], d["responsavel"], d["status"]])
    output.seek(0)
    bytes_io = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    return send_file(bytes_io, mimetype="text/csv", as_attachment=True,
                     download_name=f"relatorio_chamados_{datetime.now().strftime('%Y%m%d_%H%M')}.csv")


def _gerar_xlsx(dados, data_geracao):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Chamados"
    ws.merge_cells("A1:H1")
    ws["A1"] = "Relatório de Chamados"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:H2")
    ws["A2"] = f"Gerado em {data_geracao} • Total: {len(dados)} chamado(s)"
    ws["A2"].alignment = Alignment(horizontal="center")
    cabecalhos = ["ID", "Data", "Solicitante", "Setor Solicitante",
                  "Item", "Setor Responsável", "Responsável", "Status"]
    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(bold=True, color="FFFFFF")
    for col, cab in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=4, column=col, value=cab)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for i, d in enumerate(dados, start=5):
        ws.cell(row=i, column=1, value=d["id"])
        ws.cell(row=i, column=2, value=d["data"])
        ws.cell(row=i, column=3, value=d["solicitante"])
        ws.cell(row=i, column=4, value=d["setor_solicitante"])
        ws.cell(row=i, column=5, value=d["item"])
        ws.cell(row=i, column=6, value=d["setor_responsavel"])
        ws.cell(row=i, column=7, value=d["responsavel"])
        ws.cell(row=i, column=8, value=d["status"])
    for i, w in enumerate([6, 18, 20, 20, 40, 20, 20, 15], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True,
                     download_name=f"relatorio_chamados_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")


def _gerar_pdf(dados, data_geracao):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4),
                            leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=20)
    styles = getSampleStyleSheet()
    elementos = [
        Paragraph("<b>Relatório de Chamados</b>", styles["Title"]),
        Paragraph(f"Gerado em {data_geracao} • Total: {len(dados)} chamado(s)", styles["Normal"]),
        Spacer(1, 12),
    ]
    cabecalho = ["ID", "Data", "Solicitante", "Setor Solicitante",
                 "Item", "Setor Responsável", "Responsável", "Status"]
    linhas = [cabecalho]
    for d in dados:
        linhas.append([str(d["id"]), d["data"], d["solicitante"], d["setor_solicitante"],
                       d["item"], d["setor_responsavel"], d["responsavel"], d["status"]])
    tabela = Table(linhas, repeatRows=1, colWidths=[30, 80, 90, 80, 170, 80, 80, 60])
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F8")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    elementos.append(tabela)
    doc.build(elementos)
    output.seek(0)
    return send_file(output, mimetype="application/pdf", as_attachment=True,
                     download_name=f"relatorio_chamados_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf")


if __name__ == "__main__":
    import socket
    init_db()
    try:
        ip_local = socket.gethostbyname(socket.gethostname())
    except Exception:
        ip_local = "127.0.0.1"
    print("=" * 50)
    print("Sistema de Chamados rodando.")
    print(f"Neste computador:        http://localhost:5000")
    print(f"Outros PCs da mesma rede: http://{ip_local}:5000")
    print("=" * 50)
    app.run(host="0.0.0.0", port=5000, debug=True)